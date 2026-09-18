import logging

from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from hysprint_utils.process_specs import (
    atmospheric_args,
    field_args,
    indexed_group_args,
    optional_block_args,
)

logger = logging.getLogger(__name__)

TABLEAU_COLORS = {
    "tab:blue": "1F77B4",
    "tab:orange": "FF7F0E",
    "tab:green": "2CA02C",
    "tab:red": "D62728",
    "tab:purple": "9467BD",
    "tab:brown": "8C564B",
    "tab:pink": "E377C2",
    "tab:gray": "7F7F7F",
    "tab:olive": "BCBD22",
    "tab:cyan": "17BECF",
}
colors = list(TABLEAU_COLORS.values())


def lighten_color(hex_color, factor=0.50):
    hex_color = hex_color.lstrip("#")
    r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
    r = int(r + (255 - r) * factor)
    g = int(g + (255 - g) * factor)
    b = int(b + (255 - b) * factor)
    return f"{r:02x}{g:02x}{b:02x}".upper()


def add_experiment_sheet(workbook, process_sequence, is_testing=False):
    ws = workbook.active
    ws.title = "Experiment Data"
    start_col = 1
    incremental_number = 0

    def make_label(label, test_val=None, dropdown_key=None):
        """
        Create a label with optional test value or dropdown options.

        Args:
            label: The column header label
            test_val: Test value (if is_testing=True) or list of dropdown options
            dropdown_key: Optional key to lookup dropdown options from dropdown_options.py

        Returns:
            - Just label string if not testing and no dropdown
            - (label, test_val) tuple if testing
            - (label, test_val, dropdown_options) tuple if dropdown is specified
        """
        # Import dropdown options if needed
        dropdown_options = None
        if dropdown_key:
            try:
                from dropdown_options import get_dropdown_options

                dropdown_options = get_dropdown_options(dropdown_key)
            except ImportError:
                pass

        # If test_val is a list, treat it as dropdown options
        if isinstance(test_val, list):
            dropdown_options = test_val
            # Use first option as default test value
            default_test = test_val[0] if test_val else None
            if is_testing:
                return (label, default_test, dropdown_options)
            else:
                return (label, None, dropdown_options)

        # If dropdown_key provided, get options from config file
        if dropdown_options:
            default_test = dropdown_options[0] if dropdown_options else None
            if is_testing:
                if test_val is not None:
                    return (label, test_val, dropdown_options)
                else:
                    return (label, default_test, dropdown_options)
            else:
                return (label, None, dropdown_options)

        # Standard behavior (no dropdown)
        if is_testing:
            if test_val is not None:
                return (label, test_val)
            else:
                return (label, f"Test for {label}")
        else:
            return label

    def generate_steps_for_process(process_name, config):
        """
        This method constructs steps for each process, in the order they should
        appear as Excel columns. Field data (label + test value) comes from
        shared/hysprint_utils/process_specs.py instead of being hardcoded inline -
        see that module's docstring for why. This function still owns the actual
        assembly ORDER per process, and the handful of genuinely special-cased
        branches (Spin Coating's single-vs-multi spin step naming, Experiment
        Info's Ink Recycling variant) - real conditional logic, not duplicated
        data.
        """

        def f(key, variant="fields"):
            return make_label(*field_args(process_name, key, variant=variant))

        def indexed(group_key, count):
            steps = []
            for i in range(1, count + 1):
                for key, val in indexed_group_args(process_name, group_key, i):
                    steps.append(make_label(key, val))
            return steps

        def optional(block_key):
            return [
                make_label(key, val) for key, val in optional_block_args(process_name, block_key)
            ]

        if process_name == "Experiment Info":
            # Check if Ink Recycling exists in the sequence
            has_ink_recycling = any(p.get("process") == "Ink Recycling" for p in process_sequence)

            if has_ink_recycling:
                return [
                    f("Date", "fields_ink_recycling"),
                    f("Project_Name", "fields_ink_recycling"),
                    f("Batch", "fields_ink_recycling"),
                    f("Subbatch", "fields_ink_recycling"),
                    f("Sample", "fields_ink_recycling"),
                    f("Nomad ID", "fields_ink_recycling"),
                    f("Variation", "fields_ink_recycling"),
                ]

            return [
                # Using this format to speed up testing with nomad
                # The format is: STEP, TEST_VARIABLE
                f("Date"),
                f("Project_Name"),
                f("Batch"),
                f("Subbatch"),
                f("Sample"),
                f("Nomad ID"),
                f("Variation"),
                f("Sample dimension"),
                f("Sample area [cm^2]"),
                f("Number of pixels"),
                f("Pixel area [cm^2]"),
                f("Substrate material"),
                f("Substrate conductive layer"),
                f("Sheet Resistance [Ohms/square]"),
                f("Transmission [%]"),
                f("Number of junctions"),
                f("Notes"),
            ]

        if process_name == "Cleaning O2-Plasma" or process_name == "Cleaning UV-Ozone":
            steps = [f("Datetime"), f("Operator")]
            steps.extend(indexed("solvents", config.get("solvents", 0)))

            if process_name == "Cleaning O2-Plasma":
                steps.extend(
                    [
                        f("Gas-Plasma Gas"),
                        f("Gas-Plasma Time [s]"),
                        f("Gas-Plasma Power [W]"),
                        f("Notes"),
                    ]
                )
            if process_name == "Cleaning UV-Ozone":
                steps.extend([f("UV-Ozone Time [s]"), f("Notes")])
            return steps

        if process_name in [
            "Spin Coating",
            "Dip Coating",
            "Slot Die Coating",
            "Inkjet Printing",
            "Blade Coating",
            "Screen Printing",
        ]:
            steps = [
                f("Datetime"),
                f("Operator"),
                f("Material name"),
                f("Layer type"),
                f("Tool/GB name"),
                f("Layer thickness [nm]"),
            ]

            # Add solvent/solute steps
            steps.extend(indexed("solvents", config.get("solvents", 0)))
            steps.extend(indexed("solutes", config.get("solutes", 0)))

            steps.extend(
                [
                    f("Viscosity [mPa*s]"),
                    f("Contact angle [°]"),
                    f("Density [g/cm^3]"),
                    f("Surface tension [mN/m]"),
                ]
            )

            # Add process-specific steps
            if process_name == "Spin Coating":
                steps.extend([f("Solution volume [uL]"), f("Spin Delay [s]")])

                if config.get("spinsteps", 0) == 1:
                    steps.extend(
                        [
                            f("Rotation speed [rpm]"),
                            f("Rotation time [s]"),
                            f("Acceleration [rpm/s]"),
                        ]
                    )
                else:
                    steps.extend(indexed("spinsteps", config.get("spinsteps", 0)))

                if config.get("antisolvent", False):
                    steps.extend(optional("antisolvent"))

                if config.get("gasquenching", False):
                    steps.extend(optional("gasquenching"))

                if config.get("vacuumquenching", False):
                    steps.extend(optional("vacuumquenching"))

            elif process_name == "Slot Die Coating":
                steps.extend(
                    [
                        f("Solution volume [uL]"),
                        f("Flow rate [ul/min]"),
                        f("Head gap [mm]"),
                        f("Speed [mm/s]"),
                        f("Air knife angle [°]"),
                        f("Air knife gap [cm]"),
                        f("Bead volume [mm/s]"),
                        f("Drying speed [cm/min]"),
                        f("Chuck heating temperature [°C]"),
                    ]
                )

            elif process_name == "Dip Coating":
                steps.append(f("Dipping duration [s]"))

            elif process_name == "Blade Coating":
                steps.extend(
                    [
                        f("Solution volume [uL]"),
                        f("Blade Speed [mm/s]"),
                        f("Dispensed Ink Volume [uL]"),
                        f("Blade Gap [um]"),
                        f("Blade Size"),
                        f("Coating Width [mm]"),
                        f("Coating Length [mm]"),
                        f("Dead Length [mm]"),
                        f("Bed Temperature [°C]"),
                        f("Ink Temperature [°C]"),
                    ]
                )

                if config.get("gasquenching", False):
                    steps.extend(optional("gasquenching"))

            elif process_name == "Screen Printing":
                steps.extend(
                    [
                        f("Solution volume [uL]"),
                        f("Mesh material"),
                        f("Mesh count [meshes/cm]"),
                        f("Mesh thickness [um]"),
                        f("Thread diameter [um]"),
                        f("Mesh opening [um]"),
                        f("Mesh tension [N/cm]"),
                        f("Mesh angle [°]"),
                        f("Emulsion material"),
                        f("Emulsion thickness [um]"),
                        f("Squeegee material"),
                        f("Squeegee shape"),
                        f("Squeegee angle [°]"),
                        f("Printing speed [mm/s]"),
                        f("Printing direction"),
                        f("Printing pressure [bar]"),
                        f("Snap-off distance [mm]"),
                        f("Printing method"),
                    ]
                )

                if config.get("gasquenching", False):
                    steps.extend(optional("gasquenching"))

                if config.get("vacuumquenching", False):
                    steps.extend(optional("vacuumquenching"))

                if config.get("airknifequenching", False):
                    steps.extend(optional("airknifequenching"))

            elif process_name == "Inkjet Printing":
                steps.extend(
                    [
                        f("Printhead name"),
                        f("Number of active nozzles"),
                        f("Active nozzles"),
                        f("Droplet density X [dpi]"),
                        f("Droplet density Y [dpi]"),
                        f("Quality factor"),
                        f("Step size"),
                        f("Printing direction"),
                        f("Number of swaths"),
                        f("Printed area [mm²]"),
                        f("Droplet per second [1/s]"),
                        f("Droplet volume [pl]"),
                        f("Ink reservoir pressure [bar]"),
                        f("Table temperature [°C]"),
                        f("Dropping Height [mm]"),
                        f("Substrate thickness [mm]"),
                        f("Printing speed [mm/s]"),
                        f("Print head angle [deg]"),
                        f("Nozzle temperature [°C]"),
                        f("Nozzle voltage config file"),
                        f("Image used"),
                        # make_label('rel. humidity [%]', 45),
                    ]
                )

                if config.get("gavd", False):
                    steps.extend(optional("gavd"))

            # Add annealing steps for all coating processes
            steps.extend(
                [
                    f("Annealing time [min]"),
                    f("Annealing temperature [°C]"),
                    f("Annealing atmosphere"),
                    f("Notes"),
                ]
            )

            return steps

        # PVD Processes
        if process_name == "Evaporation" or process_name == "Sublimation":
            # "Sublimation" is a legacy alias with identical columns/test values to
            # Evaporation and (same as before this migration) no archive paths of its
            # own either way - see process_specs.py's note on the "Evaporation" entry.
            lookup_name = "Evaporation"
            return [
                make_label(*field_args(lookup_name, "Datetime")),
                make_label(*field_args(lookup_name, "Operator")),
                make_label(*field_args(lookup_name, "Material name")),
                make_label(*field_args(lookup_name, "Layer type")),
                make_label(*field_args(lookup_name, "Tool/GB name")),
                make_label(*field_args(lookup_name, "Organic")),
                make_label(*field_args(lookup_name, "Sample holder width [mm]")),
                make_label(*field_args(lookup_name, "Base pressure [bar]")),
                make_label(*field_args(lookup_name, "Pressure start [bar]")),
                make_label(*field_args(lookup_name, "Pressure end [bar]")),
                make_label(*field_args(lookup_name, "Source temperature start[°C]")),
                make_label(*field_args(lookup_name, "Source temperature end[°C]")),
                make_label(*field_args(lookup_name, "Substrate temperature [°C]")),
                make_label(*field_args(lookup_name, "Thickness [nm]")),
                make_label(*field_args(lookup_name, "Rate start [angstrom/s]")),
                make_label(*field_args(lookup_name, "Rate target [angstrom/s]")),
                make_label(*field_args(lookup_name, "Tooling factor")),
                make_label(*field_args(lookup_name, "Notes")),
            ]

        if process_name == "Co-Evaporation":
            steps = [
                f("Datetime"),
                f("Operator"),
                f("Material name"),
                f("Layer type"),
                f("Tool/GB name"),
            ]
            steps.extend(indexed("materials", config.get("materials", 0)))
            steps.append(f("Notes"))
            return steps

        if process_name == "Sputtering":
            return [
                f("Datetime"),
                f("Operator"),
                f("Material name"),
                f("Layer type"),
                f("Tool/GB name"),
                f("Gas"),
                f("Temperature [°C]"),
                f("Pressure [mbar]"),
                f("Deposition time [s]"),
                f("Burn in time [s]"),
                f("Power [W]"),
                f("Rotation rate [rpm]"),
                f("Thickness [nm]"),
                f("Gas flow rate [cm^3/min]"),
                f("Notes"),
            ]

        if process_name == "Laser Scribing":
            return [
                f("Datetime"),
                f("Operator"),
                f("Laser wavelength [nm]"),
                f("Laser pulse time [ps]"),
                f("Laser pulse frequency [kHz]"),
                f("Speed [mm/s]"),
                f("Fluence [J/cm2]"),
                f("Power [%]"),
                f("Recipe file"),
                f("Dead area [cm2]"),
                f("Width of cell [mm]"),
                f("Number of cells"),
                f("Notes"),
            ]

        if process_name == "ALD":
            return [
                f("Datetime"),
                f("Operator"),
                f("Material name"),
                f("Layer type"),
                f("Tool/GB name"),
                f("Source"),
                f("Thickness [nm]"),
                f("Temperature [°C]"),
                f("Rate [A/s]"),
                f("Time [s]"),
                f("Number of cycles"),
                f("Precursor 1"),
                f("Pulse duration 1 [s]"),
                f("Manifold temperature 1 [°C]"),
                f("Bottle temperature 1 [°C]"),
                f("Precursor 2 (Oxidizer/Reducer)"),
                f("Pulse duration 2 [s]"),
                f("Manifold temperature 2 [°C]"),
                f("Notes"),
            ]

        if process_name == "Annealing":
            return [
                f("Datetime"),
                f("Operator"),
                f("Annealing time [min]"),
                f("Annealing temperature [°C]"),
                f("Annealing athmosphere"),
                f("Relative humidity [%]"),
                f("Notes"),
            ]

        if process_name == "Generic Process":
            return [
                f("Datetime"),
                f("Operator"),
                f("Name"),
                f("Notes"),
            ]

        if process_name == "Ink Recycling":
            steps = []
            steps.extend(indexed("solvents", config.get("solvents", 0)))
            steps.extend(indexed("solutes", config.get("solutes", 0)))
            steps.extend(indexed("precursors", config.get("precursors", 0)))

            steps.extend(
                [
                    f("Functional liquid name"),
                    f("Functional liquid volume [ml]"),
                    f("Dissolving temperature [°C]"),
                ]
            )
            steps.extend(
                [
                    f("Filter material"),
                    f("Filter size [mm]"),
                    f("Filter weight [g]"),
                ]
            )
            steps.extend(
                [
                    f("Recovered solute [g]"),
                    f("Yield [%]"),
                    f("Notes"),
                ]
            )
            return steps

        else:
            logger.warning(
                "Process '%s' not defined in generate_steps_for_process. Using default steps.",
                process_name,
            )
            return [make_label("Undefined Process", "Test value")]

    for process_data in process_sequence:
        process_name = process_data["process"]
        custom_config = process_data.get("config", {})
        color_index = incremental_number % len(colors)
        cell_color = colors[color_index]
        steps = generate_steps_for_process(process_name, custom_config)

        # Append atmospheric values if requested (for all processes except Experiment Info)
        if process_name != "Experiment Info" and custom_config.get("add_atmospheric", False):
            steps.extend(make_label(key, val) for key, val in atmospheric_args())

        step_count = len(steps)
        end_col = start_col + step_count - 1

        if process_name != "Experiment Info":
            process_label = f"{incremental_number}: {process_name}"
        else:
            process_label = process_name

        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        cell = ws.cell(row=1, column=start_col)
        cell.value = process_label
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.fill = PatternFill(start_color=cell_color, end_color=cell_color, fill_type="solid")

        row2_color = lighten_color(cell_color)
        for i, step_item in enumerate(steps):
            col_index = start_col + i

            # Handle different step_item formats
            has_dropdown = False
            dropdown_options = None

            if isinstance(step_item, tuple):
                if len(step_item) == 3:
                    # (label, test_val, dropdown_options)
                    step_label, test_val, dropdown_options = step_item
                    has_dropdown = True
                elif len(step_item) == 2:
                    # (label, test_val)
                    step_label, test_val = step_item
                else:
                    step_label = step_item[0]
                    test_val = None
            else:
                # Just a label string
                step_label = step_item
                test_val = None

            # Write the label in row 2
            cell = ws.cell(row=2, column=col_index)
            cell.value = step_label
            cell.fill = PatternFill(start_color=row2_color, end_color=row2_color, fill_type="solid")

            # Write test value if testing mode
            if is_testing and test_val is not None:
                ws.cell(row=3, column=col_index, value=test_val)

            # Apply dropdown data validation if options provided
            if has_dropdown and dropdown_options:
                # Create comma-separated list for Excel validation
                options_str = ",".join([str(opt) for opt in dropdown_options])

                # Create data validation
                dv = DataValidation(type="list", formula1=f'"{options_str}"', allow_blank=True)
                dv.error = "Please select from the dropdown list"
                dv.errorTitle = "Invalid Entry"
                dv.prompt = "Select from the list"
                dv.promptTitle = "Dropdown Selection"

                # Add validation to worksheet
                ws.add_data_validation(dv)

                # Apply to entire column (rows 3-1000)
                col_letter = get_column_letter(col_index)
                dv.add(f"{col_letter}3:{col_letter}1000")
        start_col = end_col + 1
        incremental_number += 1

    # Example: Apply a custom formula for the "Nomad ID" column (example only)
    for row in range(3, 4):
        nomad_id_formula = f'=CONCATENATE("HZB_",B{row},"_",C{row},"_",D{row},"_C-",E{row})'
        ws[f"F{row}"].value = nomad_id_formula

    # Adjust column widths
    for col in ws.columns:
        max_length = 0
        column_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value and isinstance(cell.value, str):
                max_length = max(max_length, len(cell.value))
        ws.column_dimensions[column_letter].width = max_length + 2
